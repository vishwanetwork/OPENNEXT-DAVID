"""Read-only XLSX import. Export observed daily summaries, never synthetic prices."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

REGIONS = {
    'af-south-1': ('Africa (Cape Town)', 'Cape Town'),
    'ap-east-1': ('Asia Pacific (Hong Kong)', 'Hong Kong'),
    'ap-northeast-1': ('Asia Pacific (Tokyo)', 'Tokyo'),
    'ap-northeast-2': ('Asia Pacific (Seoul)', 'Seoul'),
    'ap-northeast-3': ('Asia Pacific (Osaka)', 'Osaka'),
    'ap-south-1': ('Asia Pacific (Mumbai)', 'Mumbai'),
    'ap-south-2': ('Asia Pacific (Hyderabad)', 'Hyderabad'),
    'ap-southeast-1': ('Asia Pacific (Singapore)', 'Singapore'),
    'ap-southeast-2': ('Asia Pacific (Sydney)', 'Sydney'),
    'ap-southeast-3': ('Asia Pacific (Jakarta)', 'Jakarta'),
    'ap-southeast-5': ('Asia Pacific (Malaysia)', 'Malaysia'),
    'ca-central-1': ('Canada (Central)', 'Canada Central'),
    'eu-central-1': ('Europe (Frankfurt)', 'Frankfurt'),
    'eu-central-2': ('Europe (Zurich)', 'Zurich'),
    'eu-north-1': ('Europe (Stockholm)', 'Stockholm'),
    'eu-south-1': ('Europe (Milan)', 'Milan'),
    'eu-south-2': ('Europe (Spain)', 'Spain'),
    'eu-west-1': ('Europe (Ireland)', 'Ireland'),
    'eu-west-2': ('Europe (London)', 'London'),
    'eu-west-3': ('Europe (Paris)', 'Paris'),
    'il-central-1': ('Israel (Tel Aviv)', 'Tel Aviv'),
    'me-central-1': ('Middle East (UAE)', 'UAE'),
    'sa-east-1': ('South America (Sao Paulo)', 'Sao Paulo'),
    'us-east-1': ('US East (N. Virginia)', 'N. Virginia'),
    'us-east-2': ('US East (Ohio)', 'Ohio'),
    'us-west-1': ('US West (N. California)', 'N. California'),
    'us-west-2': ('US West (Oregon)', 'Oregon'),
}

def convert(source, destination):
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    groups = {}
    seen = set()
    regions = set()
    missing = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        region, _, zone, sku, cpu, ram, raw_model, count, memory, date, avg, high, low, per_gpu = row
        if region not in REGIONS:
            raise ValueError('Add an English region mapping before import: ' + str(region))
        regions.add(region)
        day = str(date).split(' ')[0]
        identity = (region, zone, sku, day)
        if identity in seen:
            raise ValueError('Duplicate price observation at row ' + str(row_number))
        seen.add(identity)
        assert all(isinstance(v, (int, float)) and v >= 0 for v in (avg, high, low))
        assert low <= avg + .000002 and avg <= high + .000002
        model = {'T4g': 'T4G', 'RTX PRO Server 6000': 'RTX PRO 6000'}.get(raw_model, raw_model)
        if per_gpu is not None:
            assert count > 0 and abs(avg/count-per_gpu) <= .000002
        else:
            missing += 1
        key = '|'.join([region, zone, sku])
        if key not in groups:
            groups[key] = {
                'id': 'uploaded-' + hashlib.sha256(key.encode()).hexdigest()[:18],
                'provider': 'aws', 'region': region, 'zone': zone, 'instanceType': sku,
                'gpuModel': model, 'sourceGpuModel': raw_model, 'gpuCount': count,
                'gpuMemoryGb': round(memory/count, 6) if count > 0 else None,
                'instanceGpuMemoryGiB': memory, 'vCpu': cpu, 'memoryGiB': ram,
                'formFactor': None, 'term': 'Spot', 'points': [],
            }
        groups[key]['points'].append({
            'date': day, 'value': per_gpu, 'instanceHourly': avg,
            'high': round(high/count, 8) if count > 0 else None,
            'low': round(low/count, 8) if count > 0 else None,
            'instanceHigh': high, 'instanceLow': low,
        })
    book.close()
    series = sorted(groups.values(), key=lambda s: (s['gpuModel'], s['instanceType'], s['region'], s['zone']))
    for group in series:
        group['points'].sort(key=lambda p: p['date'])
    dates = sorted({p['date'] for group in series for p in group['points']})
    result = {
        'schemaVersion': 1, 'asOf': dates[-1], 'sourceKind': 'uploaded_daily_history',
        'currency': 'USD', 'unit': 'USD/GPU-hour',
        'regions': [{'code': r, 'name': REGIONS[r][0], 'city': REGIONS[r][1]} for r in sorted(regions)],
        'series': series,
        'meta': {'sourceRows': len(seen), 'seriesCount': len(series), 'regionCount': len(regions),
                 'modelCount': len({s['gpuModel'] for s in series}), 'missingPerGpuRows': missing,
                 'earliestDate': dates[0], 'latestDate': dates[-1], 'dateTimezone': 'not supplied',
                 'dailyAverageMethod': 'as supplied; weighting not specified',
                 'currencyBasis': 'USD/hour assumed from the existing Terminal; source workbook has no explicit currency metadata',
                 'sourceFilename': Path(source).name, 'sourceSha256': hashlib.sha256(Path(source).read_bytes()).hexdigest()},
    }
    text = json.dumps(result, ensure_ascii=True, separators=(',', ':'))
    assert not any('\u4e00' <= c <= '\u9fff' for c in text)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(text, encoding='utf-8')
    print(json.dumps({'file': str(destination), 'bytes': len(text), **result['meta']}, ensure_ascii=True))

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('--output', default=str(Path(__file__).resolve().parents[1] / 'data' / 'gpu-spot-history.json'))
    args = parser.parse_args()
    convert(args.source, Path(args.output))
